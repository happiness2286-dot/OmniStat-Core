import json
import os
import sys
import io
import re
import urllib.request
import datetime
import openpyxl

# Fix Windows console encoding
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
else:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

class XSMBScraper:
    """
    Primary & Fallback Web Scraper for XSMB (Xổ Số Miền Bắc) results.
    Primary web source: https://mketqua.net/
    Fallback sources: xoso.com.vn & local Excel seed (Thong_Ke_G7_Va_Top20_XSMB_2026.xlsx).
    """

    def __init__(self, excel_path="Thong_Ke_G7_Va_Top20_XSMB_2026.xlsx"):
        self.excel_path = excel_path

    def fetch_from_mketqua(self):
        """Fetch latest draw results directly from https://mketqua.net/"""
        url = "https://mketqua.net/"
        req = urllib.request.Request(
            url,
            headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
        )
        try:
            with urllib.request.urlopen(req, timeout=8) as response:
                html = response.read().decode('utf-8')
                
            # Parse Date (e.g., id="result_date">Thứ sáu ngày 04-09-2026<)
            date_match = re.search(r'id=["\']result_date["\'][^>]*>([^<]+)<', html)
            date_str = date_match.group(1).strip() if date_match else datetime.date.today().strftime("%Y-%m-%d")

            # Parse GDB (id="rs_0_0">50066<)
            gdb_match = re.search(r'id=["\']rs_0_0["\'][^>]*>(\d{5})<', html)
            if not gdb_match:
                return None

            gdb = gdb_match.group(1)
            de_2d = gdb[-2:]

            # Parse G7.1, G7.2, G7.3, G7.4
            g7_1_m = re.search(r'id=["\']rs_7_0["\'][^>]*>(\d{2})<', html)
            g7_2_m = re.search(r'id=["\']rs_7_1["\'][^>]*>(\d{2})<', html)
            g7_3_m = re.search(r'id=["\']rs_7_2["\'][^>]*>(\d{2})<', html)
            g7_4_m = re.search(r'id=["\']rs_7_3["\'][^>]*>(\d{2})<', html)

            g7_1 = g7_1_m.group(1) if g7_1_m else "00"
            g7_2 = g7_2_m.group(1) if g7_2_m else "00"
            g7_3 = g7_3_m.group(1) if g7_3_m else "00"
            g7_4 = g7_4_m.group(1) if g7_4_m else "00"

            return {
                "date": date_str,
                "gdb": gdb,
                "de_2d": de_2d,
                "head": de_2d[0],
                "tail": de_2d[1],
                "sum": (int(de_2d[0]) + int(de_2d[1])) % 10,
                "g7_1": g7_1,
                "g7_2": g7_2,
                "g7_3": g7_3,
                "g7_4": g7_4,
                "source": "mketqua.net"
            }
        except Exception as e:
            print(f"[Scraper Warning] mketqua.net fetch failed: {e}")
            return None

    def load_from_excel_seed(self):
        """Extract historical draw data from existing Excel file."""
        if not os.path.exists(self.excel_path):
            return []

        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        if "Du_Lieu_2026" not in wb.sheetnames:
            return []

        sheet = wb["Du_Lieu_2026"]
        records = []

        for r in range(2, sheet.max_row + 1):
            stt = sheet.cell(r, 1).value
            date_str = sheet.cell(r, 2).value
            gdb = sheet.cell(r, 3).value
            de_2d = sheet.cell(r, 4).value
            g7_1 = sheet.cell(r, 5).value
            g7_2 = sheet.cell(r, 8).value
            
            g7_3 = sheet.cell(r, 11).value if sheet.max_column >= 11 else None
            g7_4 = sheet.cell(r, 14).value if sheet.max_column >= 14 else None

            if date_str and gdb is not None:
                gdb_str = str(gdb).zfill(5)
                de_2d_str = str(de_2d).zfill(2) if de_2d is not None else gdb_str[-2:]
                
                records.append({
                    "stt": stt,
                    "date": str(date_str).strip(),
                    "gdb": gdb_str,
                    "de_2d": de_2d_str,
                    "head": de_2d_str[0],
                    "tail": de_2d_str[1],
                    "sum": (int(de_2d_str[0]) + int(de_2d_str[1])) % 10,
                    "g7_1": str(g7_1).zfill(2) if g7_1 is not None else "",
                    "g7_2": str(g7_2).zfill(2) if g7_2 is not None else "",
                    "g7_3": str(g7_3).zfill(2) if g7_3 is not None else "",
                    "g7_4": str(g7_4).zfill(2) if g7_4 is not None else ""
                })

        return records

    def fetch_live_xsmb(self, target_date=None):
        """
        Fetch latest draw results with primary & secondary fallback order:
        1. mketqua.net
        2. xoso.com.vn
        3. Local Excel seed
        """
        # Primary: mketqua.net
        mketqua_res = self.fetch_from_mketqua()
        if mketqua_res:
            print(f"[Scraper] Successfully fetched latest draw from mketqua.net: {mketqua_res['date']} -> GDB {mketqua_res['gdb']}")
            return mketqua_res

        # Secondary: xoso.com.vn
        try:
            url = "https://xoso.com.vn/xsmb-xo-so-mien-bac.html"
            req = urllib.request.Request(
                url,
                headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
            )
            with urllib.request.urlopen(req, timeout=5) as response:
                html = response.read().decode('utf-8')
                
            gdb_match = re.search(r'class="font28[^"]*">(\d{5})<', html)
            g7_matches = re.findall(r'class="font18[^"]*">(\d{2})<', html)
            
            if gdb_match:
                gdb = gdb_match.group(1)
                today_str = datetime.date.today().strftime("%Y-%m-%d")
                return {
                    "date": today_str,
                    "gdb": gdb,
                    "de_2d": gdb[-2:],
                    "head": gdb[-2:][0],
                    "tail": gdb[-2:][1],
                    "sum": (int(gdb[-2:][0]) + int(gdb[-2:][1])) % 10,
                    "g7_1": g7_matches[0] if len(g7_matches) > 0 else "00",
                    "g7_2": g7_matches[1] if len(g7_matches) > 1 else "00",
                    "g7_3": g7_matches[2] if len(g7_matches) > 2 else "00",
                    "g7_4": g7_matches[3] if len(g7_matches) > 3 else "00",
                    "source": "xoso.com.vn"
                }
        except Exception as e:
            print(f"[Scraper Warning] Secondary fallback failed: {e}")

        # Tertiary fallback: Excel seed data
        seed_data = self.load_from_excel_seed()
        if seed_data:
            latest = seed_data[-1]
            latest["source"] = "excel_seed"
            return latest
            
        return None


if __name__ == "__main__":
    scraper = XSMBScraper()
    live = scraper.fetch_live_xsmb()
    print("Fetched Live Data:", live)
